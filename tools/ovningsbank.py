"""Övningsbanken: Excel-fil med alla övningar i appen, till Haris.

    python tools/ovningsbank.py                 # skriver docs/Övningsbanken.xlsx
    python tools/ovningsbank.py annan/sökväg.xlsx

Hämtar allt ur data/content.json (CAP-passets 19 moment, stigens block och
tidtagningar, utbildningsvägens planerade övningar, nivåproven, avsluten i
prototyp 2) och passmallen i admin, plus taggdimensionerna i data/taggar.json
som rullistor. Haris fyller på med egna övningar och taggar; en senare
import läser tillbaka filen med samma kolumnnycklar (KOLUMNER nedan).

Färdighet, sällskap, yta och nivå på de befintliga raderna är CARLS/CLAUDES
TOLKNING av filmerna (tabellen TOLK) – Haris ska granska dem.
"""
import datetime as dt
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
PAGES = "https://8bitcat.github.io/cap-online/"
MAXRAD = 500          # rader med rullistor/format på fliken Övningar
MAXFILM = 80          # rader på fliken Filmer som Film-listan täcker

# ---------------------------------------------------------------- kolumner
# (nyckel, rubrik, bredd, grupp, slag)  slag: text|tal|x|lista:<lista>|film|jn
KOLUMNER = [
    ("id", "ID", 9, "Övning", "text"),
    ("namn", "Övning", 32, "Övning", "text"),
    ("desc", "Instruktion", 58, "Övning", "text"),
    ("typ", "Typ", 12, "Övning", "lista:typ"),
    ("tid", "Tid (min)", 8, "Övning", "tal"),
    ("maltid", "Måltid (s)", 9, "Övning", "tal"),
    ("fas", "Fas i passet", 13, "Övning", "lista:fas"),
    ("omrade", "Område", 13, "Övning", "lista:omrade"),
    ("fardighet", "Färdighet", 28, "Övning", "fri:fardighet"),
    ("niva", "Nivå (1–5)", 16, "Övning", "lista:svarighet"),
    ("film", "Film", 28, "Film", "film"),
    ("start", "Start (s)", 8, "Film", "tal"),
    ("slut", "Slut (s)", 8, "Film", "tal"),
    ("harfilm", "Finns film?", 10, "Film", "jn"),
    ("alderfran", "Ålder från", 7, "För vem", "tal"),
    ("aldertill", "Ålder till", 7, "För vem", "tal"),
    ("erfarenhet", "Erfarenhet", 16, "För vem", "fri:niva"),
    ("position", "Position", 13, "För vem", "fri:position"),
    ("u_boll", "Boll", 5, "Vad som behövs", "x"),
    ("u_koner", "Koner", 5, "Vad som behövs", "x"),
    ("u_mal", "Mål", 5, "Vad som behövs", "x"),
    ("u_vagg", "Vägg", 5, "Vad som behövs", "x"),
    ("u_medspelare", "Medspelare", 5, "Vad som behövs", "x"),
    ("u_stege", "Stege/häckar", 5, "Vad som behövs", "x"),
    ("yta", "Yta", 17, "Vad som behövs", "lista:yta"),
    ("s_ensam", "Ensam", 5, "Vad som behövs", "x"),
    ("s_foralder", "Med förälder", 5, "Vad som behövs", "x"),
    ("s_kompis", "Med kompis", 5, "Vad som behövs", "x"),
    ("plats", "Plats", 10, "Vad som behövs", "fri:plats"),
    ("kalla", "Källa i appen", 36, "Övrigt", "text"),
    ("anteckning", "Anteckning (Haris)", 30, "Övrigt", "text"),
]
KOL = {k: i + 1 for i, (k, *_r) in enumerate(KOLUMNER)}
L = lambda k: get_column_letter(KOL[k])  # kolumnbokstav
FRAN_APPEN = {"id", "namn", "desc", "film", "start", "slut", "kalla"}  # grå rubriker
UTR = {"Boll": "u_boll", "Koner": "u_koner", "Mål": "u_mal", "Vägg": "u_vagg",
       "Medspelare": "u_medspelare", "Stege/häckar": "u_stege"}
SALL = {"E": "s_ensam", "F": "s_foralder", "K": "s_kompis"}
POS = {"mv": "Målvakt", "back": "Back", "mitt": "Mittfältare", "anf": "Anfallare"}
YTA = {"L": "Liten (5×5 m)", "M": "Mellan (10×10 m)", "S": "Stor (halvplan)"}

# ------------------------------------------------- tolkningar (Haris granskar)
# f = färdigheter, s = sällskap (E ensam, F med förälder, K med kompis),
# u = utrustning utöver appens, y = yta (L/M/S), n = lägsta nivå 1–5
TOLK = {
    # CAP-passet, 19 moment (nivå enligt stigen: del 1–2 → 1, del 3–4 → 2, del 5 → 3)
    "ub1": dict(f=["Jonglering & bollkänsla", "Driv"], s="EF", n=1),
    "ub2": dict(f=["Jonglering & bollkänsla"], s="EF", n=1),
    "ub3": dict(f=["Driv", "Vändningar"], s="EF", n=1),
    "ub4": dict(f=["Snabba fötter", "Jonglering & bollkänsla"], s="EF", n=1),
    "ft1": dict(f=["Första touch"], s="FK", n=1),
    "ft2": dict(f=["Första touch"], s="FK", n=1),
    "ft3": dict(f=["Första touch"], s="FK", n=1),
    "ft4": dict(f=["Första touch"], s="FK", n=1),
    "pa1": dict(f=["Passning insida"], s="EFK", n=2),
    "pa2": dict(f=["Passning insida"], s="FK", n=2),
    "pa3": dict(f=["Passning insida", "Första touch"], s="FK", n=2),
    "dr1": dict(f=["Driv"], s="EF", n=2),
    "dr2": dict(f=["Dribbling & finter", "Vändningar"], s="EF", n=2),
    "dr3": dict(f=["Dribbling & finter"], s="EF", n=2),
    "ko1": dict(f=["Snabba fötter"], s="EF", n=3),
    "ko2": dict(f=["Reaktion", "Första touch"], s="FK", n=3),
    "ko3": dict(f=["Snabba fötter", "Driv"], s="EF", n=3),
    "ko4": dict(f=["Reaktion", "Passning insida"], s="FK", n=3),
    "ko5": dict(f=["Snabba fötter", "Balans"], s="EF", n=3),
    # Planerade övningar (film saknas) – id från utbildningsvägen eller stigen
    "e1-jong": dict(f=["Jonglering & bollkänsla"], s="E", u=["Boll"], y="L"),
    "n2-anf-avslut": dict(f=["Avslut"], s="F", u=["Boll", "Mål"]),
    "n2-back-stallning": dict(f=["Balans", "Spelförståelse"], s="F", u=["Koner"]),
    "n2-mitt-vandning": dict(f=["Vändningar", "Första touch"], s="F", u=["Boll", "Koner"]),
    "n2-mv-grepp": dict(f=["Målvaktsspel"], s="F", u=["Boll"]),
    "n3-finter": dict(f=["Dribbling & finter"], s="EF", u=["Boll", "Koner"]),
    "n3-jonglering": dict(f=["Jonglering & bollkänsla"], s="E", u=["Boll"], y="L"),
    "n3-avslut": dict(f=["Avslut"], s="F", u=["Boll", "Mål"]),
    "n3-mv-fall": dict(f=["Målvaktsspel"], s="F", u=["Boll"]),
    "n3-mv-fotter": dict(f=["Målvaktsspel", "Första touch"], s="F", u=["Boll"]),
    "n3-back-nick": dict(f=["Nickar"], s="F", u=["Boll"]),
    "n3-mitt-axel": dict(f=["Första touch", "Spelförståelse"], s="FK", u=["Boll"]),
    "n3-mitt-framat": dict(f=["Första touch", "Driv"], s="F", u=["Boll", "Koner"]),
    "n3-anf-touch": dict(f=["Första touch", "Avslut"], s="F", u=["Boll", "Mål"]),
    "n3-anf-vrist": dict(f=["Avslut"], s="F", u=["Boll", "Mål"]),
    "n4-press": dict(f=["Första touch"], s="FK", u=["Boll"]),
    "n4-langa": dict(f=["Långa tillslag"], s="FK", u=["Boll"]),
    "n4-back-inlagg": dict(f=["Långa tillslag"], s="FK", u=["Boll", "Mål"]),
    "n4-mitt-tempo": dict(f=["Långa tillslag", "Spelförståelse"], s="FK", u=["Boll", "Koner"]),
    "n4-anf-nick": dict(f=["Nickar", "Avslut"], s="F", u=["Boll", "Mål"]),
    "n5-avslut-fart": dict(f=["Avslut", "Driv"], s="F", u=["Boll", "Mål", "Koner"]),
    "n5-en-mot-en": dict(f=["Dribbling & finter", "Avslut"], s="FK", u=["Boll", "Mål"]),
    "n5-mv-straff": dict(f=["Målvaktsspel"], s="FK", u=["Boll", "Mål"]),
    "n5-anf-press": dict(f=["Avslut"], s="FK", u=["Boll", "Mål"]),
    # Tidtagningar i stigen
    "e1-slalom": dict(f=["Dribbling & finter", "Driv"], s="F", u=["Boll", "Koner"], y="M", n=1),
    "e2-vagg": dict(f=["Passning insida", "Första touch"], s="F", u=["Boll", "Vägg"], y="L", n=2),
    "e2-driv": dict(f=["Driv", "Vändningar"], s="F", u=["Boll", "Koner"], y="S", n=2),
    "e3-fotter": dict(f=["Snabba fötter"], s="F", u=["Boll"], y="L", n=3),
    "e4-konrad": dict(f=["Dribbling & finter"], s="F", u=["Boll", "Koner"], y="M", n=4),
    # Nivåprov
    "n1-prov": dict(f=["Jonglering & bollkänsla", "Första touch"], s="F", u=["Boll"], n=1),
    "n2-prov": dict(f=["Passning insida", "Dribbling & finter"], s="F", u=["Boll", "Koner"], n=2),
    "n3-prov": dict(f=["Jonglering & bollkänsla", "Dribbling & finter"], s="F", u=["Boll"], n=3),
    "n4-prov": dict(f=["Långa tillslag", "Första touch"], s="FK", u=["Boll", "Medspelare"], n=4),
    "n5-prov": dict(f=["Spelförståelse"], s="F", u=[], n=5),
    # Avsluten i prototyp 2 (Spelarkortet) och passmallen i admin
    "avslut-jonglering": dict(f=["Jonglering & bollkänsla"], s="EF", u=["Boll"], n=1),
    "avslut-straffar": dict(f=["Avslut"], s="F", u=["Boll", "Mål"], n=2),
    "avslut-en-mot-en": dict(f=["Dribbling & finter"], s="F", u=["Boll", "Koner"], n=2),
    "mall-uppvarmning": dict(f=["Driv", "Jonglering & bollkänsla"], s="EF", u=["Boll"], n=1),
    "mall-konbana": dict(f=["Dribbling & finter"], s="F", u=["Boll", "Koner"], y="M", n=1),
    "mall-spel": dict(f=[], s="FK", u=["Boll"], n=1),
}

# Passmallen i js/admin-pass.js (mall()): bara övnings- och bedömningsblocken
PASSMALL = [
    ("mall-uppvarmning", "Uppvärmning med boll", "Touch på stället, driv, vändningar.", "Övning", 10, "Uppvärmning", "Bollkontroll"),
    ("mall-konbana", "Konbana", "Fem koner, en meter mellan. Tur och retur. Tre försök, medeltiden sparas.", "Bedömning", 6, "Huvuddel", "Teknik"),
    ("mall-spel", "Spel med moment", "Använd dagens moment i spelform.", "Övning", 12, "Huvuddel", "Teknik"),
]

# ------------------------------------------------------------------- data
def las():
    c = json.loads((ROOT / "data" / "content.json").read_text(encoding="utf-8"))
    t = json.loads((ROOT / "data" / "taggar.json").read_text(encoding="utf-8"))
    return c, t


def nivanamn(c, n):
    return next(f"{x['n']} {x['name']}" for x in c["game"]["levels"] if x["n"] == n)


def rensa(desc):
    return desc.replace("Film kommer från Haris.", "").replace("Text kommer från Haris.", "").strip()


def ytaFran(space):
    s = space or ""
    return YTA["L"] if "5 × 5" in s else YTA["M"] if "10 ×" in s else ""


def rad(c, id, namn, desc, typ, kalla, tid=None, maltid=None, fas="", omrade="", film="",
        start=None, slut=None, harfilm="Nej", position="Alla", utr=(), yta="", niva=None):
    t = TOLK.get(id, {})
    r = {
        "id": id, "namn": namn, "desc": desc, "typ": typ, "tid": tid, "maltid": maltid,
        "fas": fas, "omrade": omrade, "fardighet": ", ".join(t.get("f", [])),
        "niva": nivanamn(c, t["n"]) if t.get("n") else (nivanamn(c, niva) if niva else ""),
        "film": film, "start": start, "slut": slut, "harfilm": harfilm,
        "alderfran": 7, "aldertill": 14, "erfarenhet": "Alla", "position": position,
        "yta": t.get("y") and YTA[t["y"]] or yta, "plats": "Alla", "kalla": kalla, "anteckning": "",
    }
    for u in list(utr) + list(t.get("u", [])):
        r[UTR[u]] = "x"
    for b in t.get("s", ""):
        r[SALL[b]] = "x"
    return r


def ovningar(c):
    pillar = {p["id"]: p["name"] for p in c["pillars"]}
    sess = {s["id"]: s for s in c["sessions"]}
    rader = []
    # 1. CAP-passets 19 filmade moment
    for di, sid in enumerate(c["program"]["parts"], 1):
        s = sess[sid]
        for mi, st in enumerate(s["steps"], 1):
            rader.append(rad(c, st["id"], st["name"], st["desc"], "Övning",
                             f"CAP-passet del {di} {s['title']}, moment {mi}", tid=st["minutes"],
                             fas="Uppvärmning" if sid == "uppvarmning-bollkontroll" else "Huvuddel",
                             omrade=pillar[s["pillar"]], film=s["title"], start=st["clip"]["start"],
                             slut=st["clip"]["end"], harfilm="Ja", utr=s["equipment"], yta=ytaFran(s["space"])))
    # 2. Planerade övningar: utbildningsvägen (traning utan film) + stigens block, slås ihop på titel
    plan = {}
    for lvl in c["curriculum"]["levels"]:
        poster = [(it, "Alla") for it in lvl["general"]]
        poster += [(it, POS[p]) for p, its in (lvl.get("positions") or {}).items() for it in its]
        for it, pos in poster:
            if it["kind"] == "traning" and not it.get("session"):
                kalla = f"Utbildningsvägen nivå {lvl['n']}" + (f", {pos.lower()}" if pos != "Alla" else "")
                plan[it["title"]] = rad(c, it["id"], it["title"], rensa(it["desc"]), "Övning", kalla,
                                        fas="Huvuddel", omrade=pillar[it["pillar"]], position=pos, niva=lvl["n"])
    for e in c["stig"]["etapper"]:
        for n in e["noder"]:
            if n["typ"] != "block":
                continue
            kalla = f"Stigen etapp {e['n']} ({n['id']})"
            if n["titel"] in plan:
                plan[n["titel"]]["kalla"] += " + " + kalla
            else:
                pos = ", ".join(POS[p] for p in n.get("position", [])) or "Alla"
                plan[n["titel"]] = rad(c, n["id"], n["titel"], rensa(n["desc"]), "Övning", kalla,
                                       fas="Huvuddel", omrade="Bollkontroll", position=pos, niva=e["n"])
    rader += plan.values()
    # 3. Tidtagningar i stigen
    for e in c["stig"]["etapper"]:
        for n in e["noder"]:
            if n["typ"] == "tid":
                rader.append(rad(c, n["id"], n["titel"], n["desc"], "Tidtagning",
                                 f"Stigen etapp {e['n']}, tidtagning", maltid=n["mal"], niva=e["n"]))
    # 4. Nivåprov (utbildningsvägen = stigens grindar)
    for lvl in c["curriculum"]["levels"]:
        for it in lvl["general"]:
            if it["kind"] == "prov":
                rader.append(rad(c, it["id"], it["title"], it["desc"], "Nivåprov",
                                 f"Utbildningsvägen nivå {lvl['n']} + stigens grind {lvl['n']}", niva=lvl["n"]))
    # 5. Avsluten i prototyp 2 (Spelarkortet, Planera dagens pass)
    ids = {"ub4": "avslut-jonglering", "pa2": "avslut-straffar", "dr3": "avslut-en-mot-en"}
    for cd in c["planner"]["cooldowns"]:
        rader.append(rad(c, ids[cd["step"]], cd["name"], cd["desc"], "Övning",
                         f"Prototyp 2 Spelarkortet, avslut efter {sess[cd['sessionId']]['title']}",
                         fas="Avslut", omrade=pillar[cd["pillar"]]))
    # 6. Passmallen i admin
    for id, namn, desc, typ, tid, fas, omr in PASSMALL:
        rader.append(rad(c, id, namn, desc, typ, "Passmallen i admin (Nytt pass)", tid=tid, fas=fas, omrade=omr))
    namn = [r["namn"] for r in rader]
    dubbla = {n for n in namn if namn.count(n) > 1}
    if dubbla:
        sys.exit(f"dubbla övningsnamn (uppslagen i Planera pass kräver unika): {dubbla}")
    return rader


def filmer(c):
    pillar = {p["id"]: p["name"] for p in c["pillars"]}
    ut = []
    for s in c["sessions"]:
        ut.append([s["title"], PAGES + s["video"]["local"], s["videoSeconds"], s["trainingMinutes"],
                   pillar[s["pillar"]], None, f"https://drive.google.com/file/d/{s['video']['driveId']}/view",
                   "I appen", f"Del av CAP-passet. {len(s['steps'])} moment med start/slut."])
    ut.append(["CAP – teknik, bollkontroll och koordination", "",
               None, None, "", None, "https://drive.google.com/file/d/1QRk19keshcHTluFjk6jGoonim_CUAVen/view",
               "Inte i appen", "Den långa filmen (1,4 GB) i Drive. Inte nedladdad eller klippt än – säg till om den ska in."])
    ut.append(["Exempel: Passningar (upplagd i admin)", PAGES + "media/passningar.mp4", 32, 10, "Teknik", None, "",
               "Testfilm", "Upplagd i admin som test, samma klipp som Passningar. Kan tas bort."])
    return ut


def kunskap(c):
    ut = []
    for lvl in c["curriculum"]["levels"]:
        poster = [(it, "Alla") for it in lvl["general"]]
        poster += [(it, POS[p]) for p, its in (lvl.get("positions") or {}).items() for it in its]
        for it, pos in poster:
            if it["kind"] == "kunskap":
                finns = bool(it.get("body"))
                ut.append([it["id"], it["title"], rensa(it["desc"]), nivanamn(c, lvl["n"]), pos,
                           "Finns i appen" + (" (med quiz)" if it.get("quiz") else "") if finns else "Text kommer från Haris",
                           f"Utbildningsvägen nivå {lvl['n']}"])
    return ut


# ------------------------------------------------------------------ excel
FONT = "Arial"
MORK = "1F2937"
GUL = PatternFill("solid", fgColor="FFF2CC")
GRA = PatternFill("solid", fgColor="E5E7EB")
BAND = {"Övning": "CCF5F3", "Film": "DBEAFE", "För vem": "FDE68A", "Vad som behövs": "D9F99D", "Övrigt": "F3F4F6"}
TUNN = Side(style="thin", color="D1D5DB")
RAM = Border(left=TUNN, right=TUNN, top=TUNN, bottom=TUNN)


def f(bold=False, color="000000", size=10, italic=False):
    return Font(name=FONT, bold=bold, color=color, size=size, italic=italic)


def rubrikrad(ws, rad, rubriker, gula=(), rotera=()):
    for i, r in enumerate(rubriker, 1):
        cell = ws.cell(row=rad, column=i, value=r)
        gul = r in gula
        cell.font = f(bold=True, color="000000" if gul else "FFFFFF")
        cell.fill = GUL if gul else PatternFill("solid", fgColor=MORK)
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center",
                                   text_rotation=90 if r in rotera else 0)
        cell.border = RAM


def lista_dv(ws, omr, formel, fel=True, titel="Välj i listan", text="Välj ett värde i rullistan."):
    dv = DataValidation(type="list", formula1=formel, allow_blank=True, showErrorMessage=fel,
                        errorTitle=titel, error=text, showInputMessage=False)
    ws.add_data_validation(dv)
    dv.add(omr)


def blad_ovningar(wb, c, t, rader):
    ws = wb.active
    ws.title = "Övningar"
    # Rad 1: gruppband
    kol = 1
    for grupp in dict.fromkeys(k[3] for k in KOLUMNER):
        n = sum(1 for k in KOLUMNER if k[3] == grupp)
        ws.merge_cells(start_row=1, start_column=kol, end_row=1, end_column=kol + n - 1)
        cell = ws.cell(row=1, column=kol, value=grupp)
        cell.font = f(bold=True)
        cell.fill = PatternFill("solid", fgColor=BAND[grupp])
        cell.alignment = Alignment(horizontal="center")
        kol += n
    ws.row_dimensions[1].height = 18
    # Rad 2: rubriker
    rotera = [k[1] for k in KOLUMNER if k[4] == "x"]
    gula = [k[1] for k in KOLUMNER if k[0] not in FRAN_APPEN]
    rubrikrad(ws, 2, [k[1] for k in KOLUMNER], gula=gula, rotera=rotera)
    for k in KOLUMNER:
        if k[0] in FRAN_APPEN:
            ws.cell(row=2, column=KOL[k[0]]).fill = GRA
            ws.cell(row=2, column=KOL[k[0]]).font = f(bold=True)
    ws.row_dimensions[2].height = 78
    for key, rubrik, bredd, grupp, slag in KOLUMNER:
        ws.column_dimensions[L(key)].width = bredd
    # Data
    for ri, r in enumerate(rader, 3):
        for key, rubrik, bredd, grupp, slag in KOLUMNER:
            v = r.get(key)
            cell = ws.cell(row=ri, column=KOL[key], value=v)
            cell.font = f()
            cell.border = RAM
            cell.alignment = Alignment(wrap_text=key in ("desc", "namn", "kalla", "anteckning", "fardighet"),
                                       vertical="top", horizontal="center" if slag in ("x", "tal", "jn") else "left")
            if slag == "tal":
                cell.number_format = "0"
    # Tomma rader för nya övningar: ram + format
    for ri in range(3 + len(rader), 3 + len(rader) + 60):
        for key, rubrik, bredd, grupp, slag in KOLUMNER:
            cell = ws.cell(row=ri, column=KOL[key])
            cell.font = f()
            cell.border = RAM
            if slag == "tal":
                cell.number_format = "0"
    # Rullistor och kryss (rad 3–MAXRAD)
    omr = lambda key: f"{L(key)}3:{L(key)}{MAXRAD}"
    tagg_kol = {}  # tagg-id → kolumnbokstav på fliken Taggar (fylls av blad_taggar)
    for key, rubrik, bredd, grupp, slag in KOLUMNER:
        if slag == "x":
            lista_dv(ws, omr(key), '"x"', titel="Kryss", text="Skriv x om det stämmer, annars lämna tomt.")
        elif slag == "jn":
            lista_dv(ws, omr(key), '"Ja,Nej"')
        elif slag == "film":
            lista_dv(ws, omr(key), f"'Filmer'!$A$2:$A${MAXFILM}", titel="Film",
                     text="Välj en film från fliken Filmer. Lägg till nya filmer där först.")
        elif slag.startswith("lista:") or slag.startswith("fri:"):
            tagg_kol[key] = slag.split(":")[1]
    # Villkorsformat: "Nej" i Finns film? → ljusröd
    ws.conditional_formatting.add(omr("harfilm"), CellIsRule(operator="equal", formula=['"Nej"'],
                                  fill=PatternFill("solid", fgColor="FECACA")))
    ws.freeze_panes = "C3"
    ws.auto_filter.ref = f"A2:{L('anteckning')}{2 + len(rader)}"
    return ws, tagg_kol


def blad_taggar(wb, c, t, ws_ov, tagg_kol):
    ws = wb.create_sheet("Taggar")
    listor = {
        "typ": ("Typ", ["Övning", "Tidtagning", "Bedömning", "Nivåprov"]),
    }
    for tid_, d in t.items():
        val = list(d["val"])
        if tid_ in ("niva", "position", "plats"):
            val = ["Alla"] + val
        listor[tid_] = (d["namn"], val)
    ws.cell(row=1, column=1, value="Listorna bakom rullistorna på fliken Övningar. Lägg till egna värden under en rubrik, så kommer de med i rullistan (upp till rad 40).").font = f(italic=True)
    kolumn = {}
    for i, (tid_, (namn, val)) in enumerate(listor.items(), 1):
        ws.cell(row=2, column=i, value=namn)
        for j, v in enumerate(val, 3):
            ws.cell(row=j, column=i, value=v).font = f()
        ws.column_dimensions[get_column_letter(i)].width = max(14, max(len(v) for v in val) + 2)
        kolumn[tid_] = get_column_letter(i)
    rubrikrad(ws, 2, [namn for namn, _ in listor.values()])
    ws.freeze_panes = "A3"
    # Koppla rullistorna på Övningar
    for key, rubrik, bredd, grupp, slag in KOLUMNER:
        if key not in tagg_kol:
            continue
        kl = kolumn[tagg_kol[key]]
        fri = slag.startswith("fri:")
        lista_dv(ws_ov, f"{L(key)}3:{L(key)}{MAXRAD}", f"'Taggar'!${kl}$3:${kl}$40", fel=not fri)


def blad_filmer(wb, c, ws_ov):
    ws = wb.create_sheet("Filmer")
    rubriker = ["Film", "Länk (spela upp)", "Längd (s)", "Träningstid (min)", "Område", "Antal moment i banken",
                "Original i Drive", "Status", "Anteckning"]
    rubrikrad(ws, 1, rubriker, gula=["Film", "Länk (spela upp)", "Längd (s)", "Träningstid (min)", "Område", "Anteckning"])
    ws.row_dimensions[1].height = 32
    for w, b in zip("ABCDEFGHI", [38, 52, 9, 12, 14, 12, 52, 12, 60]):
        ws.column_dimensions[w].width = b
    rader = filmer(c)
    for ri, r in enumerate(rader, 2):
        for ci, v in enumerate(r, 1):
            cell = ws.cell(row=ri, column=ci, value=v)
            cell.font = f()
            cell.border = RAM
            cell.alignment = Alignment(wrap_text=ci in (1, 9), vertical="top")
        ws.cell(row=ri, column=6, value=f"=COUNTIF('Övningar'!${L('film')}$3:${L('film')}${MAXRAD},A{ri})").alignment = Alignment(horizontal="center", vertical="top")
        if r[1]:
            ws.cell(row=ri, column=2).hyperlink = r[1]
            ws.cell(row=ri, column=2).font = f(color="0563C1")
        if r[6]:
            ws.cell(row=ri, column=7).hyperlink = r[6]
            ws.cell(row=ri, column=7).font = f(color="0563C1")
    for ri in range(2 + len(rader), MAXFILM + 1):
        for ci in range(1, 10):
            ws.cell(row=ri, column=ci).border = RAM
            ws.cell(row=ri, column=ci).font = f()
        ws.cell(row=ri, column=6, value=f'=IF(A{ri}="","",COUNTIF(\'Övningar\'!${L("film")}$3:${L("film")}${MAXRAD},A{ri}))').alignment = Alignment(horizontal="center")
    lista_dv(ws, f"E2:E{MAXFILM}", "'Taggar'!$E$3:$E$40")
    lista_dv(ws, f"H2:H{MAXFILM}", '"I appen,Ny film,Testfilm,Inte i appen"')
    ws.freeze_panes = "B2"
    return ws


def blad_kunskap(wb, c):
    ws = wb.create_sheet("Kunskap")
    rubriker = ["ID", "Rubrik", "Om vad", "Nivå", "Position", "Status", "Källa i appen", "Anteckning (Haris)"]
    rubrikrad(ws, 1, rubriker, gula=["Anteckning (Haris)"])
    ws.row_dimensions[1].height = 22
    for w, b in zip("ABCDEFGH", [16, 34, 60, 16, 13, 26, 24, 36]):
        ws.column_dimensions[w].width = b
    rader = kunskap(c)
    for ri, r in enumerate(rader, 2):
        for ci, v in enumerate(r, 1):
            cell = ws.cell(row=ri, column=ci, value=v)
            cell.font = f()
            cell.border = RAM
            cell.alignment = Alignment(wrap_text=ci in (2, 3), vertical="top")
        ws.cell(row=ri, column=8).border = RAM
    ws.conditional_formatting.add(f"F2:F{1 + len(rader)}", CellIsRule(operator="equal", formula=['"Text kommer från Haris"'],
                                  fill=PatternFill("solid", fgColor="FECACA")))
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:H{1 + len(rader)}"
    return len(rader)


def blad_planera(wb):
    ws = wb.create_sheet("Planera pass")
    ws.column_dimensions["A"].width = 5
    for w, b in zip("BCDEFG", [40, 10, 14, 30, 16, 40]):
        ws.column_dimensions[w].width = b
    ws["A1"] = "Planera ett pass"
    ws["A1"].font = f(bold=True, size=14)
    ws["A2"] = "Välj övningar i rullistan i kolumn B. Tid, fas, film och nivå hämtas från fliken Övningar. Gula celler fyller du i. Kopiera fliken (högerklicka på fliknamnet) för fler pass."
    ws["A2"].font = f(italic=True)
    for r, (etikett, varde) in enumerate([("Pass", "Pass med Haris"), ("Spelare", ""), ("Datum", ""), ("Riktmärke (min)", 55)], 4):
        ws.cell(row=r, column=1, value=etikett).font = f(bold=True)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=1)
        cell = ws.cell(row=r, column=2, value=varde)
        cell.fill = GUL
        cell.font = f(color="0000FF")
        cell.border = RAM
    ws["B6"].number_format = "yyyy-mm-dd"
    rubriker = ["#", "Övning", "Tid (min)", "Fas i passet", "Film", "Nivå (1–5)", "Anteckning"]
    rubrikrad(ws, 9, rubriker, gula=["Övning", "Anteckning"])
    ws.row_dimensions[9].height = 22
    forsta, antal = 10, 12
    upp = lambda key, r: f"IFERROR(INDEX('Övningar'!${L(key)}$3:${L(key)}${MAXRAD},MATCH($B{r},'Övningar'!${L('namn')}$3:${L('namn')}${MAXRAD},0)),\"\")"
    for i in range(antal):
        r = forsta + i
        ws.cell(row=r, column=1, value=i + 1)
        ws.cell(row=r, column=2).fill = GUL
        ws.cell(row=r, column=3, value=f'=IF($B{r}="","",{upp("tid", r)})')
        ws.cell(row=r, column=4, value=f'=IF($B{r}="","",{upp("fas", r)})')
        ws.cell(row=r, column=5, value=f'=IF($B{r}="","",{upp("film", r)})')
        ws.cell(row=r, column=6, value=f'=IF($B{r}="","",{upp("niva", r)})')
        ws.cell(row=r, column=7).fill = GUL
        for ci in range(1, 8):
            cell = ws.cell(row=r, column=ci)
            cell.font = f(color="0000FF" if ci in (2, 7) else "000000")
            cell.border = RAM
            cell.alignment = Alignment(horizontal="center" if ci in (1, 3) else "left", vertical="top", wrap_text=ci in (2, 5, 7))
        ws.cell(row=r, column=3).number_format = "0"
    sista = forsta + antal - 1
    ws.cell(row=sista + 1, column=2, value="Totalt").font = f(bold=True)
    ws.cell(row=sista + 1, column=3, value=f"=SUM(C{forsta}:C{sista})").font = f(bold=True)
    ws.cell(row=sista + 2, column=2, value="Kvar till riktmärket").font = f(bold=True)
    ws.cell(row=sista + 2, column=3, value=f"=$B$7-C{sista + 1}").font = f(bold=True)
    for r in (sista + 1, sista + 2):
        ws.cell(row=r, column=3).number_format = "0"
        ws.cell(row=r, column=3).alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=3).border = RAM
    ws.conditional_formatting.add(f"C{sista + 2}", CellIsRule(operator="lessThan", formula=["0"], font=Font(name=FONT, bold=True, color="C00000")))
    lista_dv(ws, f"B{forsta}:B{sista}", f"'Övningar'!${L('namn')}$3:${L('namn')}${MAXRAD}", titel="Övning",
             text="Välj en övning från fliken Övningar. Lägg till nya övningar där först.")
    ws.freeze_panes = "A10"
    # Exempel: de första raderna i CAP-passet, så Haris ser hur det ser ut
    for i, namn in enumerate(["Bollkänsla och första driv", "Mottagning med insidan", "Insidespassning i konrutan", "Fint och vändning i fart", "Snabba fötter vid konerna"]):
        ws.cell(row=forsta + i, column=2, value=namn)
    ws.cell(row=forsta, column=7, value="Exempelrader – byt ut mot ditt eget pass.")


def blad_lasmig(wb, c, antal, antal_kunskap):
    ws = wb.create_sheet("Läs mig", 0)
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 100
    ws["A1"] = "Övningsbanken – CAP"
    ws["A1"].font = f(bold=True, size=16)
    ws["A2"] = f"Alla övningar som finns i appen i dag ({dt.date.today():%Y-%m-%d}). Fyll på med dina egna övningar och tagga dem, så kan vi planera pass och läsa in listan i appen igen."
    ws["A2"].font = f(italic=True)
    ws["A4"], ws["B4"] = "Övningar i banken", f"=COUNTA('Övningar'!${L('namn')}$3:${L('namn')}${MAXRAD})"
    ws["A5"], ws["B5"] = "– varav med film", f"=COUNTIF('Övningar'!${L('harfilm')}$3:${L('harfilm')}${MAXRAD},\"Ja\")"
    ws["A6"], ws["B6"] = "– varav utan film (att filma)", f"=COUNTIF('Övningar'!${L('harfilm')}$3:${L('harfilm')}${MAXRAD},\"Nej\")"
    ws["A7"], ws["B7"] = "Filmer", f"=COUNTA('Filmer'!$A$2:$A${MAXFILM})"
    ws["A8"], ws["B8"] = "Kunskapstexter", f"=COUNTA('Kunskap'!$A$2:$A$200)"
    for r in range(4, 9):
        ws.cell(row=r, column=1).font = f(bold=True)
        ws.cell(row=r, column=2).font = f()
        ws.cell(row=r, column=2).alignment = Alignment(horizontal="left")
    steg = [
        ("Fliken Övningar", "En rad per övning. Raderna som redan finns kommer från appen. Lägg nya övningar på de tomma raderna längst ner och lämna ID tomt – det sätts när vi läser in filen."),
        ("Gula rubriker", "Kolumner som du fyller i eller kontrollerar. Grå rubriker kommer från appen och behöver du inte röra."),
        ("Taggarna som redan står", "Färdighet, nivå, yta och sällskap på de befintliga raderna är Carls tolkning av filmerna. Rätta det som är fel."),
        ("Rullistor", "Tryck på pilen i cellen. Listorna ligger på fliken Taggar, där kan du lägga till egna värden."),
        ("Kryss", "Kolumnerna Boll, Koner, Mål, Vägg, Medspelare, Stege/häckar och Ensam, Med förälder, Med kompis: skriv x när det stämmer, annars tomt."),
        ("Flera värden", "Färdighet, Erfarenhet, Position och Plats får innehålla flera värden med kommatecken, t.ex. \"Första touch, Passning insida\". Alla = passar alla."),
        ("Tid och måltid", "Tid (min) är hur länge momentet tar i ett pass. Måltid (s) används bara för Tidtagning: tiden barnet ska klara."),
        ("Nivå (1–5)", "Samma nivåer som i klientregistret: 1 Bollkompis, 2 Touchmästare, 3 Dribblare, 4 Spelmotor, 5 Matchvinnare. Ange den lägsta nivå övningen passar för."),
        ("Ålder", "Ålder från och till i hela år, t.ex. 7 och 10. Appen är gjord för 7–14 år."),
        ("Film", "Välj film från fliken Filmer (lägg till nya filmer där först, med länk). Start och Slut är sekunder i filmen där momentet börjar och slutar – appen hoppar dit."),
        ("Finns film?", "Nej = övningen finns i appen som text men saknar film. Röd markering visar vad som återstår att filma."),
        ("Fliken Planera pass", "Välj övningar i rullistan så räknas tiden ihop mot riktmärket. Kopiera fliken för fler pass."),
        ("Fliken Kunskap", "Texterna i appen (regler, tips, positionens jobb). Rött = texten väntar på dig."),
        ("Fliken Taggar", "Listorna bakom rullistorna. Lägg till egna värden under rätt rubrik."),
    ]
    ws["A10"] = "Så här fyller du i"
    ws["A10"].font = f(bold=True, size=12)
    for i, (a, b) in enumerate(steg, 11):
        ws.cell(row=i, column=1, value=a).font = f(bold=True)
        cell = ws.cell(row=i, column=2, value=b)
        cell.font = f()
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=i, column=1).alignment = Alignment(vertical="top")
    r = 11 + len(steg) + 1
    ws.cell(row=r, column=1, value="Exempel").font = f(bold=True, size=12)
    ws.cell(row=r + 1, column=1, value="Raden \"Bollkänsla och första driv\" överst på fliken Övningar är helt ifylld och visar hur en rad ska se ut.").font = f()
    ws.cell(row=r + 3, column=1, value="Frågor").font = f(bold=True, size=12)
    ws.cell(row=r + 4, column=1, value="Skriv i kolumnen Anteckning (Haris) på raden det gäller, eller till Carl.").font = f()
    ws.sheet_view.showGridLines = False


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    ut = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "docs" / "Övningsbanken.xlsx"
    c, t = las()
    rader = ovningar(c)
    wb = Workbook()
    ws_ov, tagg_kol = blad_ovningar(wb, c, t, rader)
    blad_taggar(wb, c, t, ws_ov, tagg_kol)
    blad_filmer(wb, c, ws_ov)
    antal_kunskap = blad_kunskap(wb, c)
    blad_planera(wb)
    blad_lasmig(wb, c, len(rader), antal_kunskap)
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    ut.parent.mkdir(parents=True, exist_ok=True)
    wb.save(ut)
    typer = {}
    for r in rader:
        typer[r["typ"]] = typer.get(r["typ"], 0) + 1
    medfilm = sum(1 for r in rader if r["harfilm"] == "Ja")
    print(f"skrev {ut}: {len(rader)} övningar ({', '.join(f'{v} {k.lower()}' for k, v in typer.items())}), "
          f"{medfilm} med film, {len(filmer(c))} filmer, {antal_kunskap} kunskapstexter")


if __name__ == "__main__":
    main()
